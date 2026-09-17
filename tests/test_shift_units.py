"""Probability-difference units across computation, persistence, and display."""

import numpy as np
import pytest
from openpyxl import load_workbook

from alphaverify.domain import scoring, shift
from alphaverify.infrastructure import artifact_io
from alphaverify.infrastructure.workspace import WorkspaceConfig
from alphaverify.presentation import workbooks


@pytest.mark.parametrize("suffix", [".npz", ".safetensors"])
@pytest.mark.parametrize("field", ["shift", "probability_shift_pp"])
def test_legacy_shift_converts_once_and_rewrites_in_new_units(tmp_path, suffix, field):
    old = tmp_path / ("old" + suffix)
    new = tmp_path / ("new" + suffix)
    values = np.array([[[-10., 20., np.nan]]])
    artifact_io._write_arrays(old, {field: values}, {"value": field})
    loaded = artifact_io.load_shift(old)
    np.testing.assert_allclose(loaded["probability_shift"], values / 100, equal_nan=True)
    assert field not in loaded
    artifact_io.save_shift(loaded, new, loaded["meta"], thin=True)
    np.testing.assert_allclose(artifact_io.load_shift(new)["probability_shift"],
                               values / 100, equal_nan=True)


@pytest.mark.parametrize("payload,meta", [
    ({"probability_shift": np.array([.1])}, {}),
    ({"probability_shift_pp": np.array([10.])}, {"shift_unit": "probability_difference"}),
    ({"shift": np.array([10.]), "probability_shift_pp": np.array([10.])}, {}),
])
def test_unknown_or_conflicting_artifact_units_are_rejected(tmp_path, payload, meta):
    path = tmp_path / "bad.npz"
    artifact_io._write_arrays(path, payload, meta)
    with pytest.raises(ValueError):
        artifact_io.load_shift(path)


def test_threshold_migration_preserves_effect_and_default():
    meta = {"asset": {"ticker": "TEST"}, "start_date": "2024-01-01"}
    old = WorkspaceConfig.from_meta({**meta, "evaluate": {"min_dev": 10}})
    new = WorkspaceConfig.from_meta({**meta, "evaluate": {
        "min_dev": .10, "shift_unit": "probability_difference",
    }})
    assert old.min_dev == new.min_dev == WorkspaceConfig.from_meta(meta).min_dev == .10
    with pytest.raises(ValueError):
        WorkspaceConfig.from_meta({**meta, "evaluate": {"shift_unit": "unknown"}})


def test_shift_storage_threshold_and_workbook_share_probability_units(tmp_path):
    cube = {
        "conditional_probability": np.array([[[.25]], [[.75]]]),
        "bin_observation_counts": np.array([[100]]),
        "bin_hit_counts": np.array([[[25]], [[75]]]),
        "eligible_observation_count": 100,
        "barriers": np.array([.01, .02]), "horizons": np.array([1]),
        "bin_edges": np.array([]), "meta": {"bin_labels": ["all"]},
    }
    result = shift.from_cube(cube, np.array([[.125], [.5]]))
    np.testing.assert_array_equal(result["probability_shift"], [[[.125]], [[.25]]])
    assert shift.evaluate(result, min_dev=.125, min_bin_n=50, min_run=2)["passed"]
    assert not shift.evaluate(result, min_dev=.126, min_bin_n=50, min_run=2)["passed"]
    path = tmp_path / "shift.xlsx"
    workbooks.write_shift_xlsx(result, path, "example")
    book = load_workbook(path)
    sheet = book.active
    assert sheet.cell(5, 2).value == .25
    assert sheet.cell(6, 2).value == .125
    assert sheet.cell(5, 2).number_format == '+0.0%;-0.0%;0.0%'
    book.close()


def test_score_rescaling_preserves_comparisons_including_ties(monkeypatch):
    # Binary-exact probabilities include a tied null and a supported zero score.
    probability = np.array([
        [[[.25], [.5]], [[.75], [.5]]],
        [[[.25], [.5]], [[.75], [.5]]],
        [[[.375], [.5]], [[.625], [.5]]],
        [[[.125], [.5]], [[.875], [.5]]],
    ])
    baseline = np.full((4, 2, 1), .5)
    counts = np.full((4, 2, 1), 100)
    new = scoring.bin_scores(probability, baseline, counts, [-.1, .1])
    original = scoring.baseline_shifts
    monkeypatch.setattr(scoring, "baseline_shifts", lambda *args: 100 * original(*args))
    old = scoring.bin_scores(probability, baseline, counts, [-.1, .1])
    np.testing.assert_allclose(new.bin_score, old.bin_score / 100)
    np.testing.assert_array_equal(new.score_supported, old.score_supported)
    np.testing.assert_array_equal(new.bin_score[1:] >= new.bin_score[0],
                                  old.bin_score[1:] >= old.bin_score[0])


def test_evaluate_can_be_restricted_to_one_bin():
    # Bin 0 has the larger run; restricting to bin 1 must report bin 1's own best cell.
    shifts = np.array([[[.30], [.12]], [[.30], [.12]], [[0.], [0.]]])
    cube = {
        "probability_shift": shifts, "conditional_probability": shifts + .5,
        "baseline_probability": np.full((3, 1), .5), "bin_hit_counts": np.full((3, 2, 1), 40),
        "bin_observation_counts": np.full((2, 1), 100),
        "barriers": np.array([-.02, -.01, .01]), "horizons": np.array([5]),
    }
    assert shift.evaluate(cube, .10, 50, 2)["best"]["bin"] == 0
    best = shift.evaluate(cube, .10, 50, 2, bins=[1])["best"]
    assert (best["bin"], best["dev"]) == (1, .12)


def test_barrier_labels_are_exact_for_the_grid_step():
    from alphaverify.infrastructure.workspace import WorkspaceConfig
    from alphaverify.presentation.display import barrier_number_format, format_barrier

    def grid(step, bound):
        return WorkspaceConfig.from_meta({"asset": {}, "start_date": "2024-01-01",
                                          "barriers": {"min": -bound, "max": bound, "step": step}}).barriers

    whole, quarter = grid(.01, .2), grid(.0025, .03)
    assert [format_barrier(value, whole) for value in (-.2, 0., .07)] == ["-20%", "+0%", "+7%"]
    assert barrier_number_format(whole) == "+0%;-0%;0%"
    labels = [format_barrier(value, quarter) for value in quarter]
    assert labels[11:14] == ["-0.25%", "+0.00%", "+0.25%"] and len(set(labels)) == len(quarter)
    assert barrier_number_format(quarter) == "+0.00%;-0.00%;0%"
    assert format_barrier(.002, grid(.002, .03)) == "+0.2%"
