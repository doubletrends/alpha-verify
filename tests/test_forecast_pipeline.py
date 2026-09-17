"""Stage 5 lists cleared nodes, combines active cleared bins one per family, and checks them against history."""

from collections import Counter
from unittest.mock import Mock

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from alphaverify.domain import barrier, combination
from alphaverify.infrastructure import artifact_io
from alphaverify.infrastructure.artifact_history import feature_bins_from_artifact
from alphaverify.infrastructure.workspace import Workspace
from alphaverify.pipeline import step_02_shift
from alphaverify.pipeline import step_04_selection as selection_stage
from alphaverify.pipeline import step_05_forecast as forecast_stage
from alphaverify.presentation import bin_figures

BARRIERS, HORIZONS = np.array([-.02, -.01, .01, .02]), np.array([1, 5])
FEATURES = {  # node: (family, feature values over time)
    "a": ("osc", lambda t: np.sin(t / 7)),
    "b": ("osc", lambda t: np.sin(t / 7 + .2)),
    "c": ("mom", lambda t: np.cos(t / 23)),
}


@pytest.fixture
def cleared_workspace(selected_workspace, monkeypatch):
    """Nodes a and b clear bins; node c is tested but never clears."""
    monkeypatch.setattr(bin_figures, "write_bin_figures", Mock(return_value=[]))
    validation = artifact_io.read_json(selected_workspace.validation_summary_path)
    template = validation["tests"][0]
    validation["tests"] += [
        {**template, "node": "b", "bin": 0, "bin_label": "x < 0.5"},
        {**template, "node": "a", "bin": 1, "bin_number": 2, "bin_label": "0.5 < x",
         "monte_carlo_p_value": .004},
        {**template, "node": "c", "family": "other", "monte_carlo_p_value": .60, "cleared": False},
    ]
    validation["cleared"] = [
        {key: value for key, value in row.items() if key != "null_scores"}
        for row in validation["tests"] if row["cleared"]
    ]
    artifact_io.write_json(selected_workspace.validation_summary_path, validation)
    selection_stage.cmd_selection(selected_workspace)
    return selected_workspace


@pytest.fixture
def forecast_workspace(tmp_path, monkeypatch):
    nodes = [{"id": "baseline", "family": "_base", "feature": "constant", "params": {}}] + [
        {"id": node, "family": family, "feature": node, "params": {}}
        for node, (family, _) in FEATURES.items()
    ]
    artifact_io.write_json(tmp_path / "workspaces" / "example" / "universe.json", {
        "meta": {"asset": {"ticker": "TEST"}, "start_date": "2024-01-01", "n_bins": 4,
                 "barriers": {"min": -.02, "max": .02, "step": .01},
                 "horizons": {"min": 1, "max": 5}},
        "families": {"all": nodes},
    })
    ws = Workspace("example", tmp_path / "workspaces")
    rng = np.random.default_rng(9)
    n = 1200
    close = 100 * np.exp(np.cumsum(rng.normal(0, .015, n)))
    data = pd.DataFrame({"open": close, "high": close * 1.008, "low": close * .992,
                         "close": close, "volume": np.ones(n)},
                        index=pd.date_range("2020-01-01", periods=n))
    t = np.arange(n)
    for node in nodes:
        values = np.zeros(n) if node["id"] == "baseline" else FEATURES[node["id"]][1](t)
        feature = pd.Series(values, index=data.index)
        edges = np.array([]) if node["id"] == "baseline" else barrier.bin_edges(feature, 4)
        cube = barrier.touch_tensor(data, feature, HORIZONS, BARRIERS, edges)
        cube.update({key: data[key].to_numpy() for key in data},
                    index=data.index.astype(str).to_numpy(), feature_values=values)
        artifact_io.save_surface(cube, ws.cube_path(node["id"]), {
            "node": node["id"], "bin_labels": barrier.bin_labels(edges),
        })
    step_02_shift.cmd_shift(ws)

    last_bins = {node: int(feature_bins_from_artifact(
        artifact_io.load_surface(ws.cube_path(node)))[-1]) for node in FEATURES}
    inactive_c = (last_bins["c"] + 1) % 4

    def selected(node, b, rank, p, q):
        return {"node": node, "family": FEATURES[node][0], "feature": node, "params": {},
                "bin": b, "bin_number": b + 1, "bin_label": f"bin {b + 1}",
                "rank": rank, "monte_carlo_p_value": p, "q_value": q}

    artifact_io.write_json(ws.validation_summary_path, {
        "tests": [{"node": node, "bin": b} for node in FEATURES for b in range(4)],
    })
    artifact_io.write_json(ws.selection_summary_path, {
        "summary": {"expected_by_chance": .6},
        "selected": [
            selected("a", last_bins["a"], 1, .001, .01),
            selected("b", last_bins["b"], 2, .002, .01),
            selected("c", last_bins["c"], 3, .003, .01),
            selected("c", inactive_c, 4, .004, .01),
        ],
    })
    monkeypatch.setattr(forecast_stage, "selection_summary_is_current", lambda ws, selection: True)
    ws.last_bins = last_bins
    return ws


def test_cleared_nodes_group_selected_bins_by_node(cleared_workspace):
    selection = artifact_io.read_json(cleared_workspace.selection_summary_path)
    tested = Counter(
        row["node"] for row in artifact_io.read_json(cleared_workspace.validation_summary_path)["tests"]
    )
    nodes = forecast_stage.cleared_nodes(selection["selected"], tested)

    assert [node["node"] for node in nodes] == ["a", "b"]
    a, b = nodes
    assert (a["bins_cleared"], a["bins_tested"], a["best_rank"]) == (2, 2, 1)
    assert [item["bin_number"] for item in a["bins"]] == [1, 2]
    assert a["best_p_value"] == .004
    assert (b["bins_cleared"], b["bins_tested"], b["best_rank"]) == (1, 2, 2)


def test_forecast_lists_cleared_nodes_and_uses_one_active_bin_per_family(forecast_workspace):
    ws = forecast_workspace
    forecast_stage.cmd_forecast(ws)
    result = artifact_io.read_json(ws.forecast_path)

    assert result["summary"] == {
        "nodes_tested": 3, "nodes_cleared": 3, "bins_tested": 12, "bins_cleared": 4,
        "expected_by_chance": pytest.approx(.6), "bins_active": 3, "conditions_used": 2,
        "nesting_violations": result["summary"]["nesting_violations"],
    }
    assert [(node["node"], node["bins_cleared"], node["bins_tested"]) for node in result["nodes"]] == [
        ("a", 1, 4), ("b", 1, 4), ("c", 2, 4),
    ]
    conditions = {row["node"]: row for row in result["conditions"]}
    assert set(conditions) == {"a", "b", "c"}
    assert conditions["a"]["used"] and conditions["c"]["used"]
    assert not conditions["b"]["used"]
    assert conditions["b"]["note"] == f"same family as a bin {ws.last_bins['a'] + 1} (rank 1)"
    assert result["as_of"] == "2023-04-14"


def test_forecast_surfaces_match_stage1_counts_and_an_independent_joint_measurement(forecast_workspace):
    ws = forecast_workspace
    forecast_stage.cmd_forecast(ws)
    result = artifact_io.read_json(ws.forecast_path)
    surface = {node: artifact_io.load_surface(ws.cube_path(node)) for node in ("baseline", "a", "c")}
    used = [("a", ws.last_bins["a"]), ("c", ws.last_bins["c"])]

    expected, expected_counts = combination.naive_bayes_probability(
        surface["baseline"]["bin_hit_counts"][:, 0, :], surface["baseline"]["bin_observation_counts"][0],
        [surface[node]["bin_hit_counts"][:, b, :] for node, b in used],
        [surface[node]["bin_observation_counts"][b] for node, b in used],
    )
    combined = np.array(result["combined_probability"], dtype=float)
    np.testing.assert_allclose(combined, expected, rtol=0, atol=1e-12)
    assert result["combined_observation_counts"] == expected_counts.astype(int).tolist()

    # Measure "every used condition held" as its own Stage 1 feature; bin 1 is the joint rate.
    both = np.all([feature_bins_from_artifact(surface[node]) == b for node, b in used], axis=0)
    data = pd.DataFrame({key: surface["baseline"][key] for key in ("open", "high", "low", "close", "volume")})
    joint_cube = barrier.touch_tensor(data, pd.Series(both.astype(float)), HORIZONS, BARRIERS, np.array([.5]))
    joint = np.array(result["joint_probability"], dtype=float)
    np.testing.assert_allclose(joint, joint_cube["conditional_probability"][:, 1, :], rtol=0, atol=1e-12)
    assert result["joint_observation_counts"] == joint_cube["bin_observation_counts"][1].tolist()
    assert np.isfinite(joint).any()


def test_forecast_workbook_tabs(forecast_workspace):
    ws = forecast_workspace
    forecast_stage.cmd_forecast(ws)
    book = load_workbook(ws.forecast_workbook_path)
    assert book.sheetnames == ["naive Bayes", "vs baseline", "historical joint",
                               "naive Bayes - joint", "conditions", "cleared nodes"]
    assert book["naive Bayes"]["A1"].value.startswith("Forecast —— as of 2023-04-14 · 2 conditions")
    assert [row[1] for row in book["conditions"].iter_rows(min_row=2, values_only=True)] == ["a", "b", "c"]
    nodes = list(book["cleared nodes"].iter_rows(values_only=True))
    assert nodes[0][:5] == ("node", "family", "feature", "cleared bins", "tested bins")
    assert [row[0] for row in nodes[1:]] == ["a", "b", "c"]
    assert nodes[3][3:5] == (2, 4)


def test_no_cleared_bins_forecasts_the_baseline(forecast_workspace):
    ws = forecast_workspace
    selection = artifact_io.read_json(ws.selection_summary_path)
    selection["selected"] = []
    artifact_io.write_json(ws.selection_summary_path, selection)
    forecast_stage.cmd_forecast(ws)
    result = artifact_io.read_json(ws.forecast_path)

    assert result["nodes"] == [] and result["conditions"] == []
    assert result["summary"]["nodes_cleared"] == 0 and result["summary"]["nodes_tested"] == 3
    np.testing.assert_allclose(
        np.array(result["combined_probability"], dtype=float),
        np.array(result["baseline_probability"], dtype=float), rtol=0, atol=1e-12,
    )
    assert len(list(load_workbook(ws.forecast_workbook_path)["cleared nodes"].iter_rows())) == 1


def test_forecast_requires_a_current_selection(selected_workspace):
    forecast_stage.cmd_forecast(selected_workspace)
    assert not selected_workspace.forecast_path.exists()
