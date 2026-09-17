"""Stage 5 groups selected bins by node and lists only nodes that cleared."""

from unittest.mock import Mock

from openpyxl import load_workbook
import pytest

from alphaverify.infrastructure import artifact_io
from alphaverify.pipeline import step_04_selection as selection_stage
from alphaverify.pipeline import step_05_summary as summary_stage
from alphaverify.pipeline.status import cmd_status
from alphaverify.presentation import bin_figures


def test_summary_lists_cleared_nodes_with_their_bins(summarized_workspace):
    summary_stage.cmd_summary(summarized_workspace)
    result = artifact_io.read_json(summarized_workspace.node_summary_path)

    assert result["summary"] == {
        "nodes_tested": 3, "nodes_cleared": 2, "bins_tested": 5, "bins_cleared": 3,
        "expected_by_chance": pytest.approx(.25),
    }
    assert [node["node"] for node in result["nodes"]] == ["a", "b"]
    a, b = result["nodes"]
    assert (a["bins_cleared"], a["bins_tested"], a["best_rank"]) == (2, 2, 1)
    assert [item["bin_number"] for item in a["bins"]] == [1, 2]
    assert a["best_p_value"] == .004
    assert (b["bins_cleared"], b["bins_tested"], b["best_rank"]) == (1, 2, 2)
    assert summary_stage.node_summary_is_current(summarized_workspace, result)

    sheet = load_workbook(summarized_workspace.node_summary_workbook_path).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:5] == ("node", "family", "feature", "cleared bins", "tested bins")
    assert [row[0] for row in rows[1:]] == ["a", "b"]
    assert rows[1][5:7] == ("1, 2", "x < 0.5; 0.5 < x")


def test_summary_requires_a_current_selection(selected_workspace):
    summary_stage.cmd_summary(selected_workspace)
    assert not selected_workspace.node_summary_path.exists()


def test_summary_becomes_stale_when_selection_changes(summarized_workspace, capsys):
    summary_stage.cmd_summary(summarized_workspace)
    result = artifact_io.read_json(summarized_workspace.node_summary_path)
    selection = artifact_io.read_json(summarized_workspace.selection_summary_path)
    selection["generated"] = "changed"
    artifact_io.write_json(summarized_workspace.selection_summary_path, selection)
    assert not summary_stage.node_summary_is_current(summarized_workspace, result)
    cmd_status(summarized_workspace)
    assert "summary: stale - run summarize" in capsys.readouterr().out


def test_no_cleared_bins_writes_an_empty_summary(selected_workspace, monkeypatch):
    monkeypatch.setattr(bin_figures, "write_bin_figures", Mock(return_value=[]))
    validation = artifact_io.read_json(selected_workspace.validation_summary_path)
    for row in validation["tests"]:
        row["cleared"] = False
    validation["cleared"] = []
    artifact_io.write_json(selected_workspace.validation_summary_path, validation)
    selection_stage.cmd_selection(selected_workspace)
    summary_stage.cmd_summary(selected_workspace)
    result = artifact_io.read_json(selected_workspace.node_summary_path)
    assert result["nodes"] == [] and result["summary"]["nodes_cleared"] == 0
    rows = list(load_workbook(selected_workspace.node_summary_workbook_path).active.iter_rows())
    assert len(rows) == 1
