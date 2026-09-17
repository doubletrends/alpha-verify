"""
The pipeline's deliverable.

Each cube workbook renders a cube as one tab per condition bin. Each tab is that bin's
whole barrier-by-horizon face, so the workbook holds every value the cube holds -- it is
a faithful view of the measurement, not a summary of it.

Stage 1 workbooks show raw conditional probabilities. Stage 2 shift workbooks show the
same full grid after subtracting the unconditional baseline. The Stage 5 workbook is the
one summary view: a row per node that cleared, listing its cleared bins.
"""

from __future__ import annotations

import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from alphaverify.presentation.display import (
    SHIFT_DISPLAY_LIMIT, barrier_number_format, feature_label,
)

_PCT_FMT = '0.0%'
_SHIFT_PCT_FMT = '+0.0%;-0.0%;0.0%'
_P_FMT = '0.0000'

_WHITE, _AMBER, _RED = 'FFFFFF', 'FFD166', 'C00000'
_BLUE = '2A78D6'
_FILL_ROW1 = PatternFill('solid', start_color='666666', end_color='666666')
_FILL_ROW2 = PatternFill('solid', start_color='B2B2B2', end_color='B2B2B2')
_FILL_NBAND = PatternFill('solid', start_color='EFEFEF', end_color='EFEFEF')
_FILL_NA = PatternFill('solid', start_color='F7F7F7', end_color='F7F7F7')
_FILL_MID = PatternFill('solid', start_color='DDDDDD', end_color='DDDDDD')
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

_COL_ROW  = 4   # horizon labels
_N_ROW    = 3   # observations behind this bin at each horizon
_DATA_ROW = 5   # first barrier row


def _sheet_names(labels: list[str]) -> list[str]:
    """
    Tab names are the conditions themselves -- 'x < 0.12', '0.12 < x < 0.34' -- so the
    tab strip reads as the ladder of bins with nothing to decode.

    Excel caps a sheet name at 31 chars and forbids : \\ / ? * [ ]. Truncation could in
    principle collide two long labels, so a numeric suffix is appended only when that
    actually happens rather than pre-emptively numbering every tab.
    """
    out: list[str] = []
    for label in labels:
        name = re.sub(r'[:\\/?*\[\]]', '-', label)[:31]
        if name in out:
            base = name[:28]
            k = 2
            while f'{base}~{k}' in out:
                k += 1
            name = f'{base}~{k}'
        out.append(name)
    return out


def _write_headers(ws, row1: str, row2: str, merge_end: str) -> None:
    rows = [
        (1, row1, Font(bold=True, size=14, color='FFFFFF'), _FILL_ROW1),
        (2, row2, Font(bold=True, size=11, color='000000'), _FILL_ROW2),
    ]
    for r, text, font, fill in rows:
        ws.merge_cells(f'A{r}:{merge_end}{r}')
        c = ws[f'A{r}']
        c.value, c.font, c.fill, c.alignment = text, font, fill, _CENTER
        ws.row_dimensions[r].height = 18


def _condition_title(node_id: str, labels: list[str], edges: np.ndarray,
                     bin_index: int, unconditional: bool) -> str:
    display_feature = feature_label(node_id)
    if unconditional:
        condition = 'all'
    elif len(edges) == 0:
        condition = labels[bin_index]
    elif bin_index == 0:
        condition = f'{display_feature} < {float(edges[0]):.3f}'
    elif bin_index == len(edges):
        condition = f'{float(edges[-1]):.3f} < {display_feature}'
    else:
        condition = f'{float(edges[bin_index - 1]):.3f} < {display_feature} < {float(edges[bin_index]):.3f}'
    return f'Condition —— {condition}'


_PROBABILITY_SCALE = dict(
    start_type='num', start_value=0,   start_color=_WHITE,
    mid_type='num',   mid_value=0.5,   mid_color=_AMBER,
    end_type='num',   end_value=1.0,   end_color=_RED,
)
_SHIFT_SCALE = dict(
    start_type='num', start_value=-SHIFT_DISPLAY_LIMIT, start_color=_BLUE,
    mid_type='num', mid_value=0, mid_color=_WHITE,
    end_type='num', end_value=SHIFT_DISPLAY_LIMIT, end_color=_RED,
)


def _write_face(
    ws,
    values: np.ndarray,
    observation_counts: np.ndarray,
    barriers: np.ndarray,
    horizons: np.ndarray,
    unit: str,
    *,
    title: str,
    subtitle: str,
    cell_value,
    number_format: str,
    color_scale: dict,
    column_width: float,
) -> None:
    """
    Write one full barrier × horizon face, shaped ``(barrier, horizon)``, onto a sheet.

    Rows run from the highest barrier at the top to the lowest at the bottom, the way a
    price ladder reads: up the sheet is up in price. The zero barrier sits in the middle and is
    shaded, marking the boundary between two different questions -- above it a cell asks
    whether the *high* reached that level, below it whether the *low* did.

    The colour scale is fixed, so flipping between tabs shows the band moving with the
    condition instead of each tab being rescaled to look alike.

    The n band under the header is the observation count behind the face at each
    horizon; it falls as t grows, because the last t bars have no realized forward
    window.
    """
    barrier_count, horizon_count = values.shape
    order = np.argsort(barriers)[::-1]        # highest barrier on the top row
    end   = get_column_letter(1 + horizon_count)
    barrier_format = barrier_number_format(barriers)
    _write_headers(ws, title, subtitle, merge_end=end)

    c = ws.cell(row=_COL_ROW, column=1, value='barrier')
    c.font, c.alignment = Font(bold=True), _CENTER
    for j, t in enumerate(horizons):
        c = ws.cell(row=_COL_ROW, column=2 + j, value=f'+{int(t)}{unit}')
        c.font, c.alignment = Font(bold=True), _CENTER

    c = ws.cell(row=_N_ROW, column=1, value='n =')
    c.font, c.fill = Font(bold=True, italic=True, size=9), _FILL_NBAND
    for j in range(horizon_count):
        c = ws.cell(row=_N_ROW, column=2 + j, value=int(observation_counts[j]))
        c.font, c.fill, c.alignment = Font(italic=True, size=9), _FILL_NBAND, _CENTER

    for r_off, i in enumerate(order):
        r = _DATA_ROW + r_off
        th = float(barriers[i])
        is_zero = abs(th) < 1e-12
        tc = ws.cell(row=r, column=1, value=th)
        tc.number_format = barrier_format
        tc.font, tc.alignment = Font(bold=True), _CENTER
        if is_zero:
            tc.fill = _FILL_MID
        for j in range(horizon_count):
            v = values[i, j]
            cell = ws.cell(row=r, column=2 + j,
                           value=None if not np.isfinite(v) else cell_value(v))
            cell.number_format = number_format
            if not np.isfinite(v):
                cell.fill = _FILL_NA

    ws.conditional_formatting.add(
        f'B{_DATA_ROW}:{end}{_DATA_ROW + barrier_count - 1}', ColorScaleRule(**color_scale)
    )

    ws.column_dimensions['A'].width = 8
    for j in range(horizon_count):
        ws.column_dimensions[get_column_letter(2 + j)].width = column_width
    ws.freeze_panes = f'B{_DATA_ROW}'


def _write_bin_faces(
    cube: dict, values: np.ndarray, path: Path, node_id: str, unit: str, **face,
) -> None:
    """One tab per condition bin; each tab is that bin's full barrier × horizon face."""
    labels = cube['meta']['bin_labels']
    bin_edges = cube['bin_edges']
    effective_bin_count = values.shape[1]
    unconditional = effective_bin_count == 1 and labels[0] == 'all'
    names = _sheet_names(list(labels))

    wb = Workbook()
    wb.remove(wb.active)
    for b in range(effective_bin_count):
        _write_face(
            wb.create_sheet(names[b]), values[:, b, :], cube['bin_observation_counts'][b],
            cube['barriers'], cube['horizons'], unit,
            title=_condition_title(node_id, labels, bin_edges, b, unconditional), **face,
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_surface_xlsx(cube: dict, path: Path, node_id: str, unit: str = 'd') -> None:
    """Stage 1 view: conditional probabilities on a fixed 0..100% scale."""
    _write_bin_faces(
        cube, cube['conditional_probability'], path, node_id, unit,
        subtitle='Conditional barrier-touch probability',
        cell_value=lambda v: round(float(v), 4),
        number_format=_PCT_FMT,
        color_scale=_PROBABILITY_SCALE,
        column_width=6.5,
    )


def write_shift_xlsx(cube: dict, path: Path, node_id: str, unit: str = 'd') -> None:
    """
    Stage 2 view: probability differences from the baseline, displayed as percentages.

    The color scale is centered at zero: blue means the barrier is touched less often
    than unconditional, red means more often.
    """
    _write_bin_faces(
        cube, cube['probability_shift'], path, node_id, unit,
        subtitle='Conditional probability minus baseline probability',
        cell_value=float,
        number_format=_SHIFT_PCT_FMT,
        color_scale=_SHIFT_SCALE,
        column_width=8.5,
    )


def _write_table(ws, columns: tuple, rows: list[tuple]) -> None:
    """Header row plus one row per record; ``columns`` holds (header, width, number format)."""
    for column, (header, width, _) in enumerate(columns, 1):
        c = ws.cell(row=1, column=column, value=header)
        c.font, c.fill, c.alignment = Font(bold=True, color='FFFFFF'), _FILL_ROW1, _CENTER
        ws.column_dimensions[get_column_letter(column)].width = width
    for row, values in enumerate(rows, 2):
        for column, (value, (_, _, number_format)) in enumerate(zip(values, columns), 1):
            c = ws.cell(row=row, column=column, value=value)
            if number_format:
                c.number_format = number_format
    ws.freeze_panes = 'B2'


_SUMMARY_COLUMNS = (
    # header, width, number format
    ('node', 24, None),
    ('family', 16, None),
    ('feature', 18, None),
    ('cleared bins', 12, '0'),
    ('tested bins', 11, '0'),
    ('bin numbers', 14, None),
    ('conditions (x = feature)', 44, None),
    ('best rank', 10, '0'),
    ('best p', 10, _P_FMT),
    ('best q', 10, _P_FMT),
)


def write_node_summary_xlsx(nodes: list[dict], path: Path) -> None:
    """Stage 5 view: one row per node that cleared, ordered by its strongest bin."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'cleared nodes'
    _write_table(ws, _SUMMARY_COLUMNS, [
        (
            node['node'], node['family'], node['feature'],
            node['bins_cleared'], node['bins_tested'],
            ', '.join(str(item['bin_number']) for item in node['bins']),
            '; '.join(item['bin_label'] for item in node['bins']),
            node['best_rank'], node['best_p_value'], node['best_q_value'],
        )
        for node in nodes
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


_CONDITION_COLUMNS = (
    ('family', 16, None),
    ('node', 24, None),
    ('bin', 6, '0'),
    ('condition (x = feature)', 30, None),
    ('latest value', 13, '0.0000'),
    ('rank', 8, '0'),
    ('p', 10, _P_FMT),
    ('q', 10, _P_FMT),
    ('used', 8, None),
    ('note', 36, None),
)


def write_forecast_xlsx(
    path: Path, *, as_of: str, barriers: np.ndarray, horizons: np.ndarray, unit: str,
    combined: np.ndarray, combined_counts: np.ndarray, baseline: np.ndarray,
    joint: np.ndarray, joint_counts: np.ndarray, conditions: list[dict],
) -> None:
    """
    Stage 6 view: the naive Bayes surface beside the historical joint rate it assumes away.

    Tabs show the combined probability, its shift from the baseline, the historical
    touch rate on bars where every contributing condition held, and their gap. Surfaces
    are ``(barrier, horizon)``; the combined n band is the weakest contributor's count.
    """
    used = sum(1 for row in conditions if row['used'])
    title = f'Forecast —— as of {as_of} · {used} condition{"" if used == 1 else "s"}'
    wb = Workbook()
    wb.remove(wb.active)
    faces = (
        ('naive Bayes', combined, combined_counts,
         'Naive Bayes barrier-touch probability (conditions treated as independent)',
         lambda v: round(float(v), 4), _PCT_FMT, _PROBABILITY_SCALE, 6.5),
        ('vs baseline', combined - baseline, combined_counts,
         'Naive Bayes probability minus baseline probability',
         float, _SHIFT_PCT_FMT, _SHIFT_SCALE, 8.5),
        ('historical joint', joint, joint_counts,
         'Historical touch rate on bars where every contributing condition held',
         lambda v: round(float(v), 4), _PCT_FMT, _PROBABILITY_SCALE, 6.5),
        ('naive Bayes - joint', combined - joint, joint_counts,
         'Naive Bayes minus historical joint rate: positive means the combination overstates',
         float, _SHIFT_PCT_FMT, _SHIFT_SCALE, 8.5),
    )
    for name, values, counts, subtitle, cell_value, number_format, scale, width in faces:
        _write_face(
            wb.create_sheet(name), values, counts, barriers, horizons, unit,
            title=title, subtitle=subtitle, cell_value=cell_value,
            number_format=number_format, color_scale=scale, column_width=width,
        )
    _write_table(wb.create_sheet('conditions'), _CONDITION_COLUMNS, [
        (
            row['family'], row['node'], row['bin_number'], row['bin_label'],
            row['latest_value'], row['rank'], row['monte_carlo_p_value'], row['q_value'],
            'yes' if row['used'] else 'no', row['note'],
        )
        for row in conditions
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
